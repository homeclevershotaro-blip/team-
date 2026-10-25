"""御見積書のExcel版 (PDFと同じ内容・同じ並び)。金額欄は数式なのでExcel上で修正すると合計も自動で変わる。"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

FONT = "游ゴシック"
GRAY = PatternFill("solid", fgColor="E6E6E6")
LIGHT = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
YEN = '#,##0;"－"#,##0;""'


def f(size=9, bold=False):
    return Font(name=FONT, size=size, bold=bold)


def put(ws, ref, value, size=9, bold=False, align="left", fill=None, wrap=False, fmt=None):
    c = ws[ref]
    c.value = value
    c.font = f(size, bold)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    return c


def box(ws, rng, fill=None):
    for row in ws[rng]:
        for c in row:
            c.border = BOX
            if fill:
                c.fill = fill


def build_xlsx(data, company, sections, out, blank=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "御見積書"
    for col, w in zip("ABCDEF", [12, 30, 7, 7, 13, 26]):
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:F1")
    put(ws, "A1", data.get("title", "御 見 積 書"), 16, True, "center")
    for c in ws[1]:
        c.border = Border(bottom=Side(style="medium"))
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A3:C3")
    put(ws, "A3", f'{data["customer"]} {data.get("honorific", "様")}', 12)
    for c in ws["A3:C3"][0]:
        c.border = Border(bottom=THIN)
    put(ws, "F3", data["date"], 9, align="right")
    ws.merge_cells("A4:F4")
    put(ws, "A4", data.get("intro", company["default_intro"]))

    # 工事情報 (左) と会社情報 (右)
    info = [("工事名称", data["project"]), ("工事場所", data["place"]), ("工期", data["period"]),
            ("お支払い条件", data["payment"]), ("有効期限", data.get("valid", company["default_valid"]))]
    for i, (k, v) in enumerate(info):
        r = 8 + i
        ws.merge_cells(f"B{r}:D{r}")
        put(ws, f"A{r}", k, 9, True, fill=GRAY)
        put(ws, f"B{r}", v, 9, wrap=True)
        box(ws, f"A{r}:D{r}")
        ws.row_dimensions[r].height = 30 if len(v) > 24 else 18
    lines = [(company["name"], 11, True), (f'{company["rep_title"]}　{company["rep_name"]}', 9, False)]
    lines += [(t, 8, False) for t in company["address_lines"] + company["license_lines"]]
    lines += [(f'担当：{data.get("staff", company["default_staff"])}', 9, False)]
    for i, (t, size, bold) in enumerate(lines):
        r = 7 + i
        ws.merge_cells(f"E{r}:F{r}")
        put(ws, f"E{r}", t, size, bold)

    # 明細
    r = 7 + len(lines) + 2
    head = r
    ws.merge_cells(f"A{r}:B{r}")
    for col, label in zip("ACDEF", ["名称・内訳", "数量", "単位", "金額（税抜）", "備考"]):
        put(ws, f"{col}{r}", label, 9, True, "center", GRAY)
    box(ws, f"A{r}:F{r}", GRAY)
    sec_cells = []
    for s in sections:
        r += 1
        sr = r
        ws.merge_cells(f"A{r}:B{r}")
        put(ws, f"A{r}", s["name"], 9, True, fill=LIGHT)
        box(ws, f"A{r}:F{r}", LIGHT)
        for it in s["items"]:
            r += 1
            ws.merge_cells(f"A{r}:B{r}")
            put(ws, f"A{r}", it["name"], 9)
            ws[f"A{r}"].alignment = Alignment(indent=1, vertical="center", wrap_text=True)
            put(ws, f"C{r}", it.get("qty", ""), 9, align="center")
            put(ws, f"D{r}", it.get("unit", ""), 9, align="center")
            put(ws, f"E{r}", None if blank else it["amount"], 9, align="right", fmt=YEN)
            put(ws, f"F{r}", it.get("note", ""), 8, wrap=True)
            ws[f"F{r}"].alignment = Alignment(indent=1, vertical="center", wrap_text=True)
            box(ws, f"A{r}:F{r}")
            if len(it.get("note", "")) > 16 or len(it["name"]) > 22:
                ws.row_dimensions[r].height = 26
        put(ws, f"E{sr}", f"=SUM(E{sr + 1}:E{r})", 9, True, "right", LIGHT, fmt=YEN)
        sec_cells.append(f"E{sr}")

    # 合計 (数式: 消費税は切り捨て、税込合計を round_unit 円単位に切り捨てて端数値引き)
    rate = company["tax_rate_percent"]
    unit = data.get("round_unit", company["round_unit"])
    rows = [("小計", "=" + "+".join(sec_cells) if sec_cells else 0, False),
            (f"消費税（{rate}%）", None, False), ("端数値引き", None, False), ("合計（税込）", None, True)]
    sub_r = r + 1
    formulas = {
        sub_r: rows[0][1],
        sub_r + 1: f"=ROUNDDOWN(E{sub_r}*{rate}/100,0)",
        sub_r + 2: f"=MOD(E{sub_r}+E{sub_r + 1},{unit})" if unit else 0,
        sub_r + 3: f"=E{sub_r}+E{sub_r + 1}-E{sub_r + 2}",
    }
    for i, (label, _, big) in enumerate(rows):
        r = sub_r + i
        ws.merge_cells(f"A{r}:D{r}")
        put(ws, f"A{r}", label, 11 if big else 9, True, "right", GRAY if big else None)
        fmt = '"－"#,##0;"－"#,##0;""' if i == 2 else YEN
        put(ws, f"E{r}", formulas[r], 11 if big else 9, True, "right", GRAY if big else None, fmt=fmt)
        box(ws, f"A{r}:F{r}", GRAY if big else None)
        if big:
            ws.row_dimensions[r].height = 22
    total_ref = f"E{r}"

    # 御見積金額 (合計と連動)
    ws.merge_cells("B6:C6")
    put(ws, "A6", "御見積金額", 9, True, "center", GRAY)
    put(ws, "B6", f"={total_ref}", 16, True, fmt='"￥"#,##0"－";;"￥"')
    put(ws, "D6", "（消費税込）", 8)
    box(ws, "A6:D6")
    ws.row_dimensions[6].height = 30

    # 注記
    r += 2
    for n in data.get("notes", []):
        ws.merge_cells(f"A{r}:F{r}")
        put(ws, f"A{r}", "※ " + n, 9, wrap=True)
        ws.row_dimensions[r].height = 26 if len(n) > 45 else 16
        r += 1

    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_view.showGridLines = False
    wb.save(out)
